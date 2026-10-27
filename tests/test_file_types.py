import unittest
import tempfile
import os
import numpy as np
import soundfile as sf
import openpyxl
import pptx
from pathlib import Path
from core.extractors import (
    format_timestamp,
    is_mostly_numbers,
    extract_rtf_file,
    extract_xlsx_file,
    extract_pptx_file,
    extract_text_file,
)
from core.searcher import HybridSearcher
from core.embedder import EmbeddingGemmaEngine


class TestFileTypeBehaviors(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_format_timestamp(self):
        self.assertEqual(format_timestamp(0), "00:00")
        self.assertEqual(format_timestamp(15), "00:15")
        self.assertEqual(format_timestamp(75), "01:15")
        self.assertEqual(format_timestamp(3665), "01:01:05")

    def test_is_mostly_numbers(self):
        self.assertTrue(is_mostly_numbers("1099, 4012.50, 99.2%, 2024, 88124, 77123"))
        self.assertTrue(is_mostly_numbers("id: 12345, val1: 9942, val2: 4412, count: 500"))
        self.assertFalse(is_mostly_numbers("This is a narrative document describing the system architecture and flow."))
        self.assertFalse(is_mostly_numbers("def calculate_total(price, tax): return price * (1 + tax)"))

    def test_category_classification(self):
        self.assertEqual(HybridSearcher.classify_extension(".pdf"), "docs")
        self.assertEqual(HybridSearcher.classify_extension(".docx"), "docs")
        self.assertEqual(HybridSearcher.classify_extension(".typ"), "docs")
        self.assertEqual(HybridSearcher.classify_extension(".rtf"), "docs")
        self.assertEqual(HybridSearcher.classify_extension(".odt"), "docs")
        self.assertEqual(HybridSearcher.classify_extension(".txt"), "docs")

        self.assertEqual(HybridSearcher.classify_extension(".xlsx"), "sheets")
        self.assertEqual(HybridSearcher.classify_extension(".ods"), "sheets")
        self.assertEqual(HybridSearcher.classify_extension(".pptx"), "sheets")
        self.assertEqual(HybridSearcher.classify_extension(".odp"), "sheets")
        self.assertEqual(HybridSearcher.classify_extension(".csv"), "sheets")

        self.assertEqual(HybridSearcher.classify_extension(".py"), "code")
        self.assertEqual(HybridSearcher.classify_extension(".rs"), "code")
        self.assertEqual(HybridSearcher.classify_extension(".ts"), "code")

        self.assertEqual(HybridSearcher.classify_extension(".jpg"), "images")
        self.assertEqual(HybridSearcher.classify_extension(".webp"), "images")
        self.assertEqual(HybridSearcher.classify_extension(".avif"), "images")
        self.assertEqual(HybridSearcher.classify_extension(".jxl"), "images")

        self.assertEqual(HybridSearcher.classify_extension(".mp4"), "videos")
        self.assertEqual(HybridSearcher.classify_extension(".mkv"), "videos")
        self.assertEqual(HybridSearcher.classify_extension(".mov"), "videos")

        self.assertEqual(HybridSearcher.classify_extension(".mp3"), "audio")
        self.assertEqual(HybridSearcher.classify_extension(".wav"), "audio")
        self.assertEqual(HybridSearcher.classify_extension(".flac"), "audio")

    def test_xlsx_extractor(self):
        xlsx_path = os.path.join(self.temp_dir.name, "budget.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Q3 Expenses"
        ws.append(["Category", "Amount", "Status"])
        ws.append(["Hosting", 450, "Approved"])
        ws.append(["Marketing", 1200, "Pending"])
        wb.save(xlsx_path)

        segments = extract_xlsx_file(xlsx_path)
        self.assertEqual(len(segments), 1)
        self.assertIn("Sheet: Q3 Expenses", segments[0].text)
        self.assertIn("Hosting", segments[0].text)
        self.assertIn("Approved", segments[0].text)

    def test_pptx_extractor(self):
        pptx_path = os.path.join(self.temp_dir.name, "pitch.pptx")
        prs = pptx.Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        title.text = "FileSift Architecture"
        subtitle.text = "Local Multimodal Search"

        notes_slide = slide.notes_slide
        notes_text_frame = notes_slide.notes_text_frame
        notes_text_frame.text = "Emphasize zero-cloud privacy and local embedding."
        prs.save(pptx_path)

        segments = extract_pptx_file(pptx_path)
        self.assertEqual(len(segments), 1)
        self.assertIn("FileSift Architecture", segments[0].text)
        self.assertIn("Speaker Notes:", segments[0].text)
        self.assertIn("Emphasize zero-cloud privacy", segments[0].text)

    def test_audio_window_slicing(self):
        wav_path = os.path.join(self.temp_dir.name, "sample.wav")
        sr = 16000
        # 45 second audio
        t = np.linspace(0, 45, 45 * sr, False)
        data = (np.sin(2 * np.pi * 440 * t) * 32767).astype(np.int16)
        sf.write(wav_path, data, sr)

        windows = EmbeddingGemmaEngine.extract_audio_windows(wav_path, window_sec=30.0, hop_sec=15.0)
        self.assertEqual(len(windows), 2)
        # First window 0 to 30
        self.assertEqual(windows[0][0], 0.0)
        self.assertEqual(windows[0][1], 30.0)
        # Second window 15 to 45
        self.assertEqual(windows[1][0], 15.0)
        self.assertEqual(windows[1][1], 45.0)


if __name__ == "__main__":
    unittest.main()
